"""Izvoz / štampa / e-mail finansijske kartice (port iz Finasije)."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _amt(v: Any) -> str:
    if v is None or v == "":
        return ""
    try:
        return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return ""


def _amt0(v: Any) -> str:
    try:
        return _amt(float(v or 0))
    except (TypeError, ValueError):
        return _amt(0)


def _saldo_dp(v: Any) -> str:
    try:
        n = float(v)
    except (TypeError, ValueError):
        return ""
    if abs(n) < 0.005:
        return _amt(0)
    side = " D" if n > 0 else " P"
    return f"{_amt(abs(n))}{side}"


def _esc(s: Any) -> str:
    return (
        str(s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _row_year(r: dict[str, Any]) -> str:
    y = str(r.get("Datum") or "").strip()[:4]
    return y if y.isdigit() else ""


def ledger_years(ledger: list[dict[str, Any]]) -> list[str]:
    years = {
        str(r.get("Datum") or "")[:4]
        for r in ledger
        if str(r.get("Datum") or "").strip()[:4].isdigit()
    }
    return sorted(years, reverse=True)


def filter_ledger_year(
    ledger: list[dict[str, Any]], year: str | None
) -> list[dict[str, Any]]:
    """Kartica za jednu godinu: početno stanje + promet te godine."""
    if not year:
        return list(ledger)
    y = str(year)
    prior = [r for r in ledger if _row_year(r) and _row_year(r) < y]
    year_rows = [r for r in ledger if _row_year(r) == y]
    opening = 0.0
    if prior:
        try:
            opening = round(float(prior[-1].get("Saldo") or 0), 2)
        except (TypeError, ValueError):
            opening = 0.0
    dug_open = opening if opening > 0.0001 else None
    potr_open = abs(opening) if opening < -0.0001 else None
    out: list[dict[str, Any]] = [
        {
            "Datum": f"{y}-01-01",
            "Vrsta": "Početno stanje",
            "Dokument": "",
            "Opis": f"Saldo na 01.01.{y}",
            "Duguje": dug_open,
            "Potražuje": potr_open,
            "Saldo": opening,
            "invoice_id": None,
        }
    ]
    saldo = opening
    for r in year_rows:
        dug = float(r.get("Duguje") or 0)
        potr = float(r.get("Potražuje") or 0)
        saldo = round(saldo + dug - potr, 2)
        item = dict(r)
        item["Saldo"] = saldo
        out.append(item)
    return out


def promet_summary(rows: list[dict[str, Any]]) -> dict[str, float]:
    if not rows:
        return {
            "open_d": 0.0,
            "open_p": 0.0,
            "tek_d": 0.0,
            "tek_p": 0.0,
            "uk_d": 0.0,
            "uk_p": 0.0,
            "saldo": 0.0,
        }
    first = rows[0]
    is_open = str(first.get("Vrsta") or "") == "Početno stanje"
    open_d = float(first.get("Duguje") or 0) if is_open else 0.0
    open_p = float(first.get("Potražuje") or 0) if is_open else 0.0
    rest = rows[1:] if is_open else rows
    tek_d = round(sum(float(r.get("Duguje") or 0) for r in rest), 2)
    tek_p = round(sum(float(r.get("Potražuje") or 0) for r in rest), 2)
    try:
        saldo = round(float(rows[-1].get("Saldo") or 0), 2)
    except (TypeError, ValueError):
        saldo = 0.0
    return {
        "open_d": round(open_d, 2),
        "open_p": round(open_p, 2),
        "tek_d": tek_d,
        "tek_p": tek_p,
        "uk_d": round(open_d + tek_d, 2),
        "uk_p": round(open_p + tek_p, 2),
        "saldo": saldo,
    }


def promet_zbir_table_html(tot: dict[str, float]) -> str:
    return (
        '<table class="zbir">'
        "<thead><tr>"
        "<th rowspan='2'></th>"
        "<th colspan='2'>Početno stanje</th>"
        "<th colspan='2'>Tekući promet</th>"
        "<th colspan='2'>Ukupan promet</th>"
        "<th rowspan='2' class='num'>Saldo</th>"
        "</tr><tr>"
        "<th class='num'>D</th><th class='num'>P</th>"
        "<th class='num'>D</th><th class='num'>P</th>"
        "<th class='num'>D</th><th class='num'>P</th>"
        "</tr></thead><tbody><tr>"
        "<td><b>UKUPNO:</b></td>"
        f"<td class='num'>{_esc(_amt0(tot['open_d']))}</td>"
        f"<td class='num'>{_esc(_amt0(tot['open_p']))}</td>"
        f"<td class='num'>{_esc(_amt0(tot['tek_d']))}</td>"
        f"<td class='num'>{_esc(_amt0(tot['tek_p']))}</td>"
        f"<td class='num'>{_esc(_amt0(tot['uk_d']))}</td>"
        f"<td class='num'>{_esc(_amt0(tot['uk_p']))}</td>"
        f"<td class='num'><b>{_esc(_saldo_dp(tot['saldo']))}</b></td>"
        "</tr></tbody></table>"
    )


def ledger_rows_clean(ledger: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for r in ledger:
        out.append(
            {
                "Datum": r.get("Datum") or "",
                "Vrsta": r.get("Vrsta") or "",
                "Dokument": r.get("Dokument") or "",
                "Opis": r.get("Opis") or "",
                "Duguje": r.get("Duguje"),
                "Potražuje": r.get("Potražuje"),
                "Saldo": r.get("Saldo"),
            }
        )
    return out


def build_kartica_xlsx(
    *,
    naziv: str,
    pib: str,
    ledger: list[dict[str, Any]],
    year: str | None = None,
    brand: str = "ProRačun",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kartica"
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    bold = Font(bold=True, size=12)
    hdr = Font(bold=True, size=9)
    fill = PatternFill("solid", fgColor="F3F3F3")
    center = Alignment(horizontal="center", vertical="center")

    tot = promet_summary(ledger)
    ws["A1"] = naziv
    ws["A1"].font = bold
    period = f"PROMET {year}" if year else "PROMET"
    ws["A2"] = f"PIB {pib} · {period} · {brand} · D = duguje, P = potražuje"
    ws["A3"] = (
        f"Početno D {_amt0(tot['open_d'])} / P {_amt0(tot['open_p'])}"
        f"  ·  Tekući D {_amt0(tot['tek_d'])} / P {_amt0(tot['tek_p'])}"
        f"  ·  Ukupan D {_amt0(tot['uk_d'])} / P {_amt0(tot['uk_p'])}"
        f"  ·  Saldo {_saldo_dp(tot['saldo'])}"
    )

    start = 5
    for col, text in ((1, "Datum"), (2, "Vrsta"), (3, "Dokument"), (4, "Opis"), (7, "Saldo")):
        ws.merge_cells(start_row=start, start_column=col, end_row=start + 1, end_column=col)
        cell = ws.cell(start, col, text)
        cell.font = hdr
        cell.fill = fill
        cell.alignment = center
        cell.border = thin
        ws.cell(start + 1, col).border = thin
        ws.cell(start + 1, col).fill = fill
    ws.merge_cells(start_row=start, start_column=5, end_row=start, end_column=6)
    promet_cell = ws.cell(start, 5, "PROMET")
    promet_cell.font = hdr
    promet_cell.fill = fill
    promet_cell.alignment = center
    promet_cell.border = thin
    ws.cell(start, 6).border = thin
    ws.cell(start, 6).fill = fill
    for col, text in ((5, "D"), (6, "P")):
        cell = ws.cell(start + 1, col, text)
        cell.font = hdr
        cell.fill = fill
        cell.alignment = center
        cell.border = thin

    for i, r in enumerate(ledger_rows_clean(ledger), start + 2):
        is_open = r["Vrsta"] == "Početno stanje"
        vals = [
            r["Datum"],
            r["Vrsta"],
            r["Dokument"],
            r["Opis"],
            r["Duguje"] if r["Duguje"] is not None else None,
            r["Potražuje"] if r["Potražuje"] is not None else None,
            _saldo_dp(r["Saldo"]) if r["Saldo"] is not None else "",
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i, col, v)
            cell.border = thin
            if is_open:
                cell.font = Font(bold=True)
            if col in (5, 6) and isinstance(v, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
            if col == 7:
                cell.alignment = Alignment(horizontal="right")

    for i, w in enumerate([12, 16, 22, 45, 12, 12, 14], 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_kartica_html(
    *,
    naziv: str,
    pib: str,
    summary: dict[str, Any] | None = None,
    ledger: list[dict[str, Any]],
    auto_print: bool = False,
    for_email: bool = False,
    year: str | None = None,
    brand: str = "ProRačun",
) -> str:
    tot = promet_summary(ledger)
    last_day = ""
    for r in reversed(ledger):
        if str(r.get("Vrsta") or "") != "Početno stanje" and r.get("Datum"):
            last_day = str(r.get("Datum") or "")
            break
    period = ""
    if year:
        period = f"PROMET {year} ZA PERIOD: 01.01.{year}"
        if last_day:
            period += f" — {last_day}"
    rows = []
    for r in ledger_rows_clean(ledger):
        is_open = r["Vrsta"] == "Početno stanje"
        tr_style = ' style="background:#f7f7f7;font-weight:600;"' if is_open else ""
        rows.append(
            f"<tr{tr_style}>"
            f"<td>{_esc(r['Datum'])}</td>"
            f"<td>{_esc(r['Vrsta'])}</td>"
            f"<td>{_esc(r['Dokument'])}</td>"
            f"<td>{_esc(r['Opis'])}</td>"
            f"<td class='num'>{_esc(_amt(r['Duguje']))}</td>"
            f"<td class='num'>{_esc(_amt(r['Potražuje']))}</td>"
            f"<td class='num'><b>{_esc(_saldo_dp(r['Saldo']))}</b></td>"
            "</tr>"
        )
    onload = ' onload="window.print()"' if auto_print and not for_email else ""
    print_btn = (
        ""
        if for_email
        else '<button class="noprint" onclick="window.print()" style="margin-bottom:12px;padding:8px 14px;">Štampaj</button>'
    )
    year_bit = f" · {year}" if year else ""
    zbir = "" if for_email else promet_zbir_table_html(tot)
    period_html = "" if for_email or not period else f"<br/>{_esc(period)}"
    summary = summary or {}
    summary_html = ""
    if summary and not for_email:
        summary_html = (
            f"<div class='meta'>Fakturisano {_esc(_amt(summary.get('fakturisano')))} € · "
            f"Uplaćeno {_esc(_amt(summary.get('uplaceno')))} € · "
            f"<b>Dug {_esc(_amt(summary.get('dug')))} €</b></div>"
        )
    return f"""<!DOCTYPE html>
<html lang="sr">
<head>
<meta charset="utf-8"/>
<title>Kartica — {_esc(naziv)}{year_bit}</title>
<style>
  @page {{ size: A4; margin: 12mm; }}
  body {{ font-family: Arial, Helvetica, sans-serif; font-size: 11px; color: #111; margin: 16px; }}
  h1 {{ font-size: 18px; margin: 0 0 4px; }}
  .meta {{ color: #555; margin-bottom: 10px; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ border-bottom: 1px solid #ddd; padding: 5px 6px; vertical-align: top; }}
  th {{ text-align: left; background: #f3f3f3; font-size: 10px; text-transform: uppercase; letter-spacing: .02em; }}
  .num {{ text-align: right; white-space: nowrap; font-variant-numeric: tabular-nums; }}
  table.zbir th {{ text-align: center; }}
  table.zbir {{ margin-bottom: 16px; }}
  .foot {{ margin-top: 14px; color: #555; }}
  @media print {{ .noprint {{ display: none !important; }} }}
</style>
</head>
<body{onload}>
  {print_btn}
  <h1>{_esc(naziv)}</h1>
  <div class="meta">PIB {_esc(pib)} · {_esc(brand)}{year_bit}{period_html}</div>
  {summary_html}
  {zbir}
  <table>
    <thead>
      <tr>
        <th rowspan="2">Datum</th>
        <th rowspan="2">Vrsta</th>
        <th rowspan="2">Dokument</th>
        <th rowspan="2">Opis</th>
        <th colspan="2" class="num">Promet</th>
        <th rowspan="2" class="num">Saldo</th>
      </tr>
      <tr>
        <th class="num">D</th>
        <th class="num">P</th>
      </tr>
    </thead>
    <tbody>
      {''.join(rows) or '<tr><td colspan="7">Nema stavki</td></tr>'}
    </tbody>
  </table>
  <div class="foot">D = duguje · P = potražuje.</div>
</body>
</html>
"""


def wrap_mail_html(*, intro: str, kartica_html: str) -> str:
    """Poruka iznad tabele kartice u istom HTML tijelu."""
    paras = [p.strip() for p in (intro or "").split("\n\n")]
    chunks: list[str] = []
    for p in paras:
        if not p:
            continue
        chunks.append(
            "<p style='margin:0 0 12px;white-space:pre-wrap;font-family:Arial,Helvetica,sans-serif;font-size:14px;color:#111;'>"
            + _esc(p).replace("\n", "<br/>")
            + "</p>"
        )
    intro_html = "".join(chunks)
    idx = kartica_html.lower().find("<body")
    if idx < 0:
        return intro_html + kartica_html
    gt = kartica_html.find(">", idx)
    if gt < 0:
        return intro_html + kartica_html
    return kartica_html[: gt + 1] + intro_html + kartica_html[gt + 1 :]
