"""Izvoz liste faktura (Excel / PDF HTML)."""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _amt(v: Any) -> float | None:
    if v is None or v == "":
        return None
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def _esc(s: Any) -> str:
    return (
        str(s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def build_invoices_xlsx(
    *,
    rows: list[dict[str, Any]],
    tenant_name: str,
    tenant_pib: str,
    headers: dict[str, str],
    brand: str = "ProRačun",
) -> bytes:
    """rows: number, status, doc, client, date, amount (float|None), pay."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Fakture"

    bold = Font(bold=True, size=12)
    hdr = Font(bold=True, size=10)
    fill = PatternFill("solid", fgColor="F0F0F0")
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    right = Alignment(horizontal="right")

    ws["A1"] = tenant_name
    ws["A1"].font = bold
    ws["A2"] = f"PIB {tenant_pib} · {brand}"
    ws["A3"] = headers.get("subtitle", "")

    cols = (
        ("number", 18),
        ("status", 14),
        ("doc", 20),
        ("client", 32),
        ("date", 18),
        ("amount", 14),
        ("pay", 16),
    )
    start = 5
    for i, (key, width) in enumerate(cols, 1):
        cell = ws.cell(start, i, headers.get(key, key))
        cell.font = hdr
        cell.fill = fill
        cell.border = thin
        ws.column_dimensions[chr(64 + i)].width = width

    total = 0.0
    for r_i, row in enumerate(rows, start + 1):
        vals = [
            row.get("number") or "",
            row.get("status") or "",
            row.get("doc") or "",
            row.get("client") or "",
            row.get("date") or "",
            _amt(row.get("amount")),
            row.get("pay") or "",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r_i, c, v if v is not None else "")
            cell.border = thin
            if c == 6 and isinstance(v, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = right
                total += float(v)

    sum_row = start + 1 + len(rows)
    ws.cell(sum_row, 5, headers.get("total", "Ukupno")).font = Font(bold=True)
    sum_cell = ws.cell(sum_row, 6, round(total, 2))
    sum_cell.font = Font(bold=True)
    sum_cell.number_format = "#,##0.00"
    sum_cell.alignment = right

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_invoices_pdf_html(
    *,
    rows: list[dict[str, Any]],
    tenant_name: str,
    tenant_pib: str,
    headers: dict[str, str],
    brand: str = "ProRačun",
    auto_print: bool = True,
) -> str:
    body_rows = []
    total = 0.0
    for row in rows:
        amt = _amt(row.get("amount"))
        if amt is not None:
            total += amt
        amt_txt = f"{amt:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if amt is not None else "—"
        body_rows.append(
            "<tr>"
            f"<td>{_esc(row.get('number'))}</td>"
            f"<td>{_esc(row.get('status'))}</td>"
            f"<td>{_esc(row.get('doc'))}</td>"
            f"<td>{_esc(row.get('client'))}</td>"
            f"<td>{_esc(row.get('date'))}</td>"
            f"<td class='num'>{_esc(amt_txt)} €</td>"
            f"<td>{_esc(row.get('pay'))}</td>"
            "</tr>"
        )
    total_txt = f"{total:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    onload = ' onload="window.print()"' if auto_print else ""
    print_lbl = _esc(headers.get("print", "Štampaj / PDF"))
    subtitle = _esc(headers.get("subtitle", ""))
    return f"""<!DOCTYPE html>
<html lang="sr-ME">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_esc(headers.get("title", "Fakture"))}</title>
<style>
  body {{ font-family: system-ui, Segoe UI, sans-serif; color: #111; margin: 1.5rem; font-size: 13px; }}
  h1 {{ font-size: 1.25rem; margin: 0 0 0.25rem; }}
  .meta {{ color: #555; margin-bottom: 1rem; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ border-bottom: 1px solid #ddd; padding: 0.4rem 0.5rem; text-align: left; vertical-align: top; }}
  th {{ background: #f4f4f4; font-size: 0.75rem; text-transform: uppercase; letter-spacing: 0.03em; }}
  .num {{ text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }}
  tfoot td {{ font-weight: 700; border-top: 2px solid #222; }}
  .noprint {{ margin-bottom: 12px; padding: 8px 14px; cursor: pointer; }}
  @media print {{ .noprint {{ display: none; }} body {{ margin: 0; }} }}
</style>
</head>
<body{onload}>
<button type="button" class="noprint" onclick="window.print()">{print_lbl}</button>
<h1>{_esc(tenant_name)}</h1>
<div class="meta">PIB {_esc(tenant_pib)} · {_esc(brand)}{" · " + subtitle if subtitle else ""}</div>
<table>
<thead>
<tr>
  <th>{_esc(headers.get("number", "Broj"))}</th>
  <th>{_esc(headers.get("status", "Status"))}</th>
  <th>{_esc(headers.get("doc", "Tip"))}</th>
  <th>{_esc(headers.get("client", "Klijent"))}</th>
  <th>{_esc(headers.get("date", "Datum"))}</th>
  <th class="num">{_esc(headers.get("amount", "Ukupno"))}</th>
  <th>{_esc(headers.get("pay", "Plaćanje"))}</th>
</tr>
</thead>
<tbody>
{"".join(body_rows) if body_rows else f'<tr><td colspan="7">{_esc(headers.get("empty", "Nema faktura."))}</td></tr>'}
</tbody>
<tfoot>
<tr>
  <td colspan="5">{_esc(headers.get("total", "Ukupno"))} ({len(rows)})</td>
  <td class="num">{_esc(total_txt)} €</td>
  <td></td>
</tr>
</tfoot>
</table>
</body>
</html>
"""
